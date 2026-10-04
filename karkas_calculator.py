"""Локальный калькулятор стоимости каркасного дома.

Настольное приложение (tkinter). Прайс загружается пользователем
из файла .xlsx (лист «Прайс») или обновляется с публичной папки
Яндекс.Диска (кнопка «Обновить данные»). Все данные обрабатываются
только на компьютере пользователя.
"""

import configparser
import ctypes
import datetime
import json
import os
import re
import socket
import sys
import tempfile
import threading
import urllib.error
import urllib.parse
import urllib.request
from collections import namedtuple
from itertools import groupby

import tkinter as tk
import tkinter.font as tkfont
from tkinter import filedialog, messagebox

try:
    import openpyxl
    from openpyxl import Workbook
    from openpyxl.styles import Font as XlFont
    from openpyxl.utils import get_column_letter
except ImportError:
    openpyxl = None

FONT = "Segoe UI"
BG = "#0e0e15"
BG_SOFT = "#14141d"
CARD = "#1a1a26"
CARD_HI = "#242434"
SEL_ROW = "#231e3b"
LINE = "#272739"
ACCENT = "#7b6bff"
ACCENT_HI = "#9a8dff"
GREEN = "#37d67e"
GREEN_HI = "#45e08c"
GREEN_DK = "#1f9e5c"
AMBER = "#ffab4a"
TEXT = "#eceaf6"
MUTED = "#8f8ea9"
FAINT = "#6b6a85"
IND_OFF = "#3a3a55"

SINGLE_CATEGORIES = {"фундамент", "каркас и стены", "кровля"}

CONFIG_PATH = os.path.join(os.path.dirname(os.path.abspath(__file__)), "config.ini")


def get_or_prompt_yandex_key(parent=None):
    default_key = "https://disk.yandex.ru/i/P5pXisDjb2EIBA"
    config = configparser.ConfigParser()
    has_config = os.path.exists(CONFIG_PATH)
    val = ""
    if has_config:
        try:
            config.read(CONFIG_PATH, encoding="utf-8")
            if config.has_section("Settings"):
                val = config.get("Settings", "YA_PUBLIC_KEY", fallback="").strip()
        except Exception:
            pass
    if not has_config or not val or val == "YOUR_TOKEN_HERE":
        from tkinter import simpledialog
        entered = simpledialog.askstring(
            "Яндекс.Диск — Настройка",
            "Введите публичную ссылку или токен Яндекс.Диска для прайса:",
            initialvalue=(val if val else default_key),
            parent=parent
        )
        if entered is not None and entered.strip():
            val = entered.strip()
        else:
            val = default_key if not val else val
        try:
            if not config.has_section("Settings"):
                config.add_section("Settings")
            config.set("Settings", "YA_PUBLIC_KEY", val)
            with open(CONFIG_PATH, "w", encoding="utf-8") as f:
                config.write(f)
        except Exception:
            pass
    return val


YA_API = "https://cloud-api.yandex.net/v1/disk/public/resources"
YA_TMP_PATH = os.path.join(tempfile.gettempdir(), "karkas_price_yandex.xlsx")

Item = namedtuple("Item", "cat name unit price comment area_based")


def lerp_color(c1, c2, t):
    a = tuple(int(c1[i:i + 2], 16) for i in (1, 3, 5))
    b = tuple(int(c2[i:i + 2], 16) for i in (1, 3, 5))
    return "#%02x%02x%02x" % tuple(round(a[i] + (b[i] - a[i]) * t) for i in range(3))


def money(v):
    return f"{round(v):,}".replace(",", " ") + " ₽"


def parse_number(value):
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    s = str(value).strip().replace("\xa0", "").replace(" ", "").replace("₽", "")
    if not s:
        return None
    s = s.replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return None


def norm_header(value):
    if value is None:
        return ""
    return re.sub(r"[\s.\-()]", "", str(value).strip().lower())


def parse_price(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    try:
        ws = None
        for name in wb.sheetnames:
            if name.strip().lower() == "прайс":
                ws = wb[name]
                break
        if ws is None:
            ws = wb[wb.sheetnames[0]]
        raw = [list(r) for r in ws.iter_rows(values_only=True)]
        sheet = ws.title
    finally:
        wb.close()

    header_idx = None
    col = {}
    for i, row in enumerate(raw[:30]):
        m = {}
        for j, cell in enumerate(row):
            n = norm_header(cell)
            if not n:
                continue
            if n.startswith("категор"):
                m["cat"] = j
            elif n.startswith("позиц") or n.startswith("наимен"):
                m["name"] = j
            elif n.startswith("ед"):
                m["unit"] = j
            elif n.startswith("цен"):
                m["price"] = j
            elif n.startswith("коммент") or n.startswith("примеч"):
                m["comment"] = j
        if "cat" in m and "name" in m and "price" in m:
            header_idx = i
            col = m
            break
    if header_idx is None:
        raise ValueError("не найдена строка заголовков (Категория, Позиция, Цена)")

    items = []
    skipped = 0
    last_cat = "Прочее"
    for row in raw[header_idx + 1:]:
        def get(key, _row=row):
            j = col.get(key)
            if j is None or j >= len(_row):
                return None
            return _row[j]

        catv = get("cat")
        if catv is not None and str(catv).strip():
            last_cat = str(catv).strip()
        namev = get("name")
        name = str(namev).strip() if namev is not None else ""
        if not name:
            continue
        price = parse_number(get("price"))
        if price is None:
            skipped += 1
            continue
        unit = get("unit")
        unit_s = str(unit).strip() if unit is not None else ""
        comment = get("comment")
        comment_s = str(comment).strip() if comment is not None else ""
        unit_n = unit_s.lower().replace("²", "2")
        area_based = "м2" in unit_n or "м^2" in unit_n
        items.append(Item(last_cat, name, unit_s, price, comment_s, area_based))
    return items, skipped, sheet


def _ya_json(url, timeout=20):
    req = urllib.request.Request(url, headers={
        "User-Agent": "KarkasCalculator/1.0",
        "Accept": "application/json"})
    with urllib.request.urlopen(req, timeout=timeout) as resp:
        body = resp.read().decode("utf-8")
    try:
        data = json.loads(body)
    except ValueError:
        raise ValueError("Яндекс.Диск вернул некорректный ответ")
    if isinstance(data, dict) and data.get("error"):
        raise ValueError("Яндекс.Диск: " + str(
            data.get("description") or data.get("message") or data.get("error")))
    return data


def _ya_download(href, dest, timeout=120):
    req = urllib.request.Request(href, headers={"User-Agent": "KarkasCalculator/1.0"})
    tmp = dest + ".part"
    try:
        with urllib.request.urlopen(req, timeout=timeout) as resp, open(tmp, "wb") as f:
            while True:
                chunk = resp.read(1 << 16)
                if not chunk:
                    break
                f.write(chunk)
        os.replace(tmp, dest)
    finally:
        if os.path.exists(tmp):
            try:
                os.remove(tmp)
            except OSError:
                pass


def _fmt_ya_date(value):
    if not value:
        return "неизвестно"
    s = str(value)
    try:
        dt = datetime.datetime.fromisoformat(s.replace("Z", "+00:00"))
    except ValueError:
        return s[:16]
    try:
        dt = dt.astimezone()
    except (OSError, OverflowError, ValueError):
        pass
    return dt.strftime("%d.%m.%Y %H:%M")


def fetch_yandex_price(parent=None):
    pub_key = get_or_prompt_yandex_key(parent)
    key = urllib.parse.quote(pub_key, safe="")
    meta = _ya_json(f"{YA_API}?public_key={key}&limit=200")
    emb = meta.get("_embedded") or {}
    items = emb.get("items") or []
    candidates = [it for it in items if str(it.get("name", "")).lower().endswith(".xlsx")]
    if candidates:
        file_meta = max(candidates, key=lambda it: str(it.get("modified") or ""))
    elif str(meta.get("name", "")).lower().endswith(".xlsx"):
        file_meta = meta
    else:
        raise ValueError("В публичной папке Яндекс.Диска не найден файл .xlsx с прайсом")
    if file_meta is meta:
        dl = _ya_json(f"{YA_API}/download?public_key={key}")
    else:
        rel = file_meta.get("path") or "/" + str(file_meta.get("name", ""))
        dl = _ya_json(f"{YA_API}/download?public_key={key}&path={urllib.parse.quote(rel, safe='/')}")
    href = dl.get("href")
    if not href:
        raise ValueError("Яндекс.Диск не вернул ссылку на скачивание")
    _ya_download(href, YA_TMP_PATH)
    return YA_TMP_PATH, str(file_meta.get("name") or "прайс.xlsx"), str(file_meta.get("modified") or "")


def diff_items(old, new):
    def key(it):
        return (it.cat.strip().lower(), it.name.strip().lower())
    om = {key(it): it for it in old}
    nm = {key(it): it for it in new}
    added = [it for k, it in nm.items() if k not in om]
    removed = [it for k, it in om.items() if k not in nm]
    changed = [(om[k], it) for k, it in nm.items()
               if k in om and abs(om[k].price - it.price) > 0.005]
    return added, removed, changed


def friendly_net_error(e):
    if isinstance(e, urllib.error.HTTPError):
        if e.code == 404:
            return "Публичный файл или папка Яндекс.Диска не найдены (404) — проверьте ссылку и доступ"
        return f"Яндекс.Диск вернул ошибку HTTP {e.code}"
    if isinstance(e, urllib.error.URLError):
        r = getattr(e, "reason", e)
        s = str(r).lower()
        if isinstance(r, (socket.timeout, TimeoutError)) or "timed out" in s or "timeout" in s:
            return "Яндекс.Диск не отвечает (превышено время ожидания)"
        if isinstance(r, socket.gaierror) or "getaddrinfo" in s or "no such host" in s or "temporary failure" in s:
            return "Нет подключения к интернету — не удалось определить адрес Яндекс.Диска"
        return f"Не удалось соединиться с Яндекс.Диском ({r})"
    if isinstance(e, (socket.timeout, TimeoutError)):
        return "Яндекс.Диск не отвечает (превышено время ожидания)"
    if isinstance(e, ConnectionError):
        return "Соединение с Яндекс.Диском прервано"
    if isinstance(e, ValueError):
        return str(e)
    return f"Ошибка обновления: {e}"


class FlatButton(tk.Button):
    def __init__(self, master, base, hover, **kw):
        kw.setdefault("relief", "flat")
        kw.setdefault("bd", 0)
        kw.setdefault("highlightthickness", 0)
        kw.setdefault("cursor", "hand2")
        kw.setdefault("activebackground", hover)
        kw.setdefault("bg", base)
        super().__init__(master, **kw)
        self._base = base
        self._hover = hover
        self._job = None
        self.bind("<Enter>", lambda e: self._fade(self._hover), add="+")
        self.bind("<Leave>", lambda e: self._fade(self._base), add="+")
        self.bind("<ButtonPress-1>", self._press, add="+")
        self.bind("<ButtonRelease-1>", lambda e: self._fade(self._hover), add="+")

    def _press(self, _):
        try:
            self.config(bg=lerp_color(self._hover, "#ffffff", 0.28))
        except tk.TclError:
            pass

    def _fade(self, target):
        if self._job:
            try:
                self.after_cancel(self._job)
            except tk.TclError:
                return
            self._job = None
        try:
            start = self["bg"]
        except tk.TclError:
            return
        if start == target:
            return
        steps = 5
        state = [0]

        def tick():
            try:
                state[0] += 1
                self.config(bg=lerp_color(start, target, state[0] / steps))
                if state[0] < steps:
                    self._job = self.after(16, tick)
            except tk.TclError:
                return

        tick()


class ScrollFrame(tk.Frame):
    _wheel_bound = False

    def __init__(self, master, bg=BG):
        super().__init__(master, bg=bg)
        self.canvas = tk.Canvas(self, bg=bg, highlightthickness=0, bd=0)
        self.vsb = tk.Scrollbar(self, orient="vertical", command=self.canvas.yview,
                                width=10, troughcolor=bg, bg=CARD_HI,
                                activebackground=ACCENT, elementborderwidth=0)
        self.inner = tk.Frame(self.canvas, bg=bg)
        self._win = self.canvas.create_window((0, 0), window=self.inner, anchor="nw")
        self.inner.bind("<Configure>", lambda e: self.canvas.configure(scrollregion=self.canvas.bbox("all")))
        self.canvas.bind("<Configure>", self._on_canvas_conf)
        self.canvas.configure(yscrollcommand=self.vsb.set)
        self.canvas.pack(side="left", fill="both", expand=True)
        self.vsb.pack(side="right", fill="y")
        if not ScrollFrame._wheel_bound:
            self.bind_all("<MouseWheel>", ScrollFrame._on_global_wheel)
            ScrollFrame._wheel_bound = True

    @staticmethod
    def _on_global_wheel(event):
        try:
            w = event.widget.winfo_containing(event.x_root, event.y_root)
        except (AttributeError, KeyError, tk.TclError):
            return
        if w is None or isinstance(w, tk.Scrollbar):
            return
        while w is not None and not isinstance(w, ScrollFrame):
            w = getattr(w, "master", None)
        if w is None:
            return
        if not event.delta:
            return
        w.canvas.yview_scroll(-int(event.delta / 120), "units")

    def _on_canvas_conf(self, e):
        self.canvas.itemconfig(self._win, width=e.width)


class OptionRow(tk.Frame):
    def __init__(self, master, mode, item, on_change):
        super().__init__(master, bg=CARD, padx=10, pady=8)
        self.mode = mode
        self.item = item
        self.on_change = on_change
        self.selected = False
        self.hover = False
        self._t = 0.0
        self._job = None

        self.canvas = tk.Canvas(self, width=18, height=18, bg=CARD,
                                highlightthickness=0, cursor="hand2")
        col = tk.Frame(self, bg=CARD)
        self.name_lbl = tk.Label(col, text=item.name, bg=CARD, fg=TEXT,
                                 font=(FONT, 10), anchor="w", justify="left",
                                 wraplength=420, cursor="hand2")
        self.name_lbl.pack(fill="x")
        unit_disp = item.unit or ("м²" if item.area_based else "шт")
        sub = " · ".join(x for x in (unit_disp, item.comment) if x)
        self.sub_lbl = None
        if sub:
            self.sub_lbl = tk.Label(col, text=sub, bg=CARD, fg=MUTED,
                                    font=(FONT, 8), anchor="w", justify="left",
                                    wraplength=420, cursor="hand2")
            self.sub_lbl.pack(fill="x")
        self.cost_lbl = tk.Label(self, text=money(item.price) + ("/м²" if item.area_based else ""),
                                 bg=CARD, fg=MUTED, font=(FONT, 9, "bold"), cursor="hand2")

        self.canvas.pack(side="left")
        self.cost_lbl.pack(side="right", padx=(6, 2))
        col.pack(side="left", fill="x", expand=True, padx=(10, 6))

        self._widgets = [self, self.canvas, col, self.name_lbl, self.cost_lbl]
        if self.sub_lbl:
            self._widgets.append(self.sub_lbl)
        for w in self._widgets:
            w.bind("<Button-1>", self._click)
        self.bind("<Enter>", self._enter, add="+")
        self.bind("<Leave>", self._leave, add="+")
        self._draw()

    def _bg(self):
        return SEL_ROW if self.selected else (CARD_HI if self.hover else CARD)

    def _apply_bg(self):
        bg = self._bg()
        for w in self._widgets:
            w.config(bg=bg)
        self.cost_lbl.config(fg=ACCENT_HI if self.selected else MUTED)
        self._draw()

    def _enter(self, _):
        self.hover = True
        self._apply_bg()

    def _leave(self, _):
        try:
            x, y = self.winfo_pointerxy()
            rx, ry = self.winfo_rootx(), self.winfo_rooty()
        except tk.TclError:
            return
        if rx <= x < rx + self.winfo_width() and ry <= y < ry + self.winfo_height():
            return
        self.hover = False
        self._apply_bg()

    def _click(self, _):
        if self.mode == "single":
            if not self.selected:
                self.on_change(self)
        else:
            self.set_selected(not self.selected)
            self.on_change(self)

    def set_selected(self, value, animate=True):
        if self.selected == value:
            self._apply_bg()
            return
        self.selected = value
        self._apply_bg()
        target = 1.0 if value else 0.0
        if not animate:
            self._t = target
            self._draw()
            return
        if self._job:
            try:
                self.after_cancel(self._job)
            except tk.TclError:
                return
        start = self._t
        steps = 6
        state = [0]

        def tick():
            try:
                state[0] += 1
                self._t = start + (target - start) * state[0] / steps
                self._draw()
                if state[0] < steps:
                    self._job = self.after(16, tick)
            except tk.TclError:
                return

        tick()

    def _draw(self):
        try:
            c = self.canvas
        except AttributeError:
            return
        c.delete("all")
        s = 18
        m = 2.5
        col = lerp_color(IND_OFF, ACCENT, self._t)
        if self.mode == "single":
            c.create_oval(m, m, s - m, s - m, outline=col, width=2)
            r = 4.5 * self._t
            if r > 0.7:
                c.create_oval(9 - r, 9 - r, 9 + r, 9 + r, fill=col, outline="")
        else:
            c.create_rectangle(m, m, s - m, s - m, outline=col, width=2,
                               fill=lerp_color(self._bg(), ACCENT, self._t))
            if self._t > 0.5:
                c.create_line(5, 9, 8, 12, 13, 6, fill="#ffffff", width=2,
                              capstyle="round", joinstyle="round")


class App(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("Калькулятор стоимости каркасного дома")
        self.geometry("1160x760")
        self.minsize(1000, 640)
        self.configure(bg=BG)
        self.items = []
        self.file_path = None
        self.calculated = False
        self.dirty = False
        self.last_rows = []
        self.last_total = 0.0
        self.last_area = 100
        self._shown = 0
        self._total_job = None
        self.blocks = {}
        self._syncing = False
        self._sync_result = None
        self._local_dir = None
        self._head_narrow = False
        self._head_last_w = 0
        self._status_narrow = False
        self._status_last_w = 0
        self.area_var = tk.IntVar(value=100)
        self._build()
        self.rebuild_categories()

    def _build(self):
        head = tk.Frame(self, bg=BG_SOFT, highlightbackground=LINE, highlightthickness=1, padx=18, pady=10)
        head.pack(fill="x")
        self.head = head
        self.head_title_lbl = tk.Label(head, text="Калькулятор стоимости каркасного дома",
                                       bg=BG_SOFT, fg=TEXT, font=(FONT, 15, "bold"))
        self.head_sub_lbl = tk.Label(head, text="расчёты выполняются локально, данные не покидают компьютер",
                                     bg=BG_SOFT, fg=FAINT, font=(FONT, 9))
        self.head_btns = tk.Frame(head, bg=BG_SOFT)
        self.btn_sync = FlatButton(self.head_btns, base=ACCENT, hover=ACCENT_HI, text="Обновить данные",
                                   command=self.sync_yandex, fg="#ffffff", font=(FONT, 10, "bold"),
                                   padx=18, pady=9)
        self.btn_sync.pack(side="left")
        self.btn_load = FlatButton(self.head_btns, base=CARD_HI, hover="#2c2c46", text="Загрузить из файла",
                                   command=self.load_dialog, fg=TEXT, font=(FONT, 10),
                                   padx=18, pady=9)
        self.btn_load.pack(side="left", padx=(10, 0))
        self.head_title_lbl.grid(row=0, column=0, sticky="w")
        self.head_sub_lbl.grid(row=1, column=0, sticky="w", pady=(1, 0))
        self.head_btns.grid(row=0, column=1, rowspan=2, sticky="e")
        head.grid_columnconfigure(0, weight=1)
        head.bind("<Configure>", self._on_head_resize, add="+")

        self.area_card = tk.Frame(self, bg=CARD, highlightbackground=LINE, highlightthickness=1, padx=18, pady=12)
        self.area_card.pack(fill="x", padx=16, pady=(12, 12))
        row = tk.Frame(self.area_card, bg=CARD)
        row.pack(fill="x")
        tk.Label(row, text="Площадь дома", bg=CARD, fg=TEXT, font=(FONT, 11, "bold")).pack(side="left")
        self.area_val = tk.Label(row, text="100 м²", bg=CARD, fg=ACCENT_HI, font=(FONT, 14, "bold"))
        self.area_val.pack(side="right")
        tk.Scale(self.area_card, from_=30, to=300, orient="horizontal", variable=self.area_var,
                 showvalue=False, resolution=1, bg=CARD, troughcolor=BG,
                 highlightthickness=0, bd=0, activebackground=ACCENT_HI,
                 sliderrelief="flat", width=12, sliderlength=24,
                 command=self._on_area).pack(fill="x", pady=(6, 0))
        ends = tk.Frame(self.area_card, bg=CARD)
        ends.pack(fill="x")
        tk.Label(ends, text="30 м²", bg=CARD, fg=FAINT, font=(FONT, 8)).pack(side="left")
        tk.Label(ends, text="300 м²", bg=CARD, fg=FAINT, font=(FONT, 8)).pack(side="right")

        self.sync_card = tk.Frame(self, bg=CARD, highlightbackground=LINE, highlightthickness=1, padx=18, pady=10)
        tk.Label(self.sync_card, text="ДАННЫЕ С ЯНДЕКС.ДИСКА", bg=CARD, fg=MUTED,
                 font=(FONT, 8, "bold"), anchor="w").pack(fill="x")
        self.sync_body = tk.Frame(self.sync_card, bg=CARD, height=72)
        self.sync_body.pack(fill="x", pady=(4, 0))
        self.sync_body.pack_propagate(False)
        self.sync_sf = ScrollFrame(self.sync_body, bg=CARD)
        self.sync_sf.pack(fill="both", expand=True)
        self.sync_info_lbl = tk.Label(self.sync_sf.inner, text="", bg=CARD, fg=TEXT,
                                      font=(FONT, 9), anchor="w", justify="left")
        self.sync_info_lbl.pack(fill="x")
        self.sync_sf.canvas.bind("<Configure>", self._on_sync_canvas_conf, add="+")

        self.status = tk.Frame(self, bg=BG_SOFT, highlightbackground=LINE, highlightthickness=1, padx=16, pady=8)
        self.status.pack(fill="x", side="bottom")
        self.file_lbl = tk.Label(self.status, text="Файл не загружен", bg=BG_SOFT, fg=FAINT, font=(FONT, 8))
        self.note_lbl = tk.Label(self.status, text="Загрузите прайс — «Обновить данные» или «Загрузить из файла»",
                                 bg=BG_SOFT, fg=AMBER, font=(FONT, 9))
        self.file_lbl.grid(row=0, column=0, sticky="w")
        self.note_lbl.grid(row=0, column=1, sticky="e")
        self.status.grid_columnconfigure(0, weight=1)
        self.status.bind("<Configure>", self._on_status_resize, add="+")

        bar = tk.Frame(self, bg=BG)
        bar.pack(fill="x", side="bottom", padx=16, pady=(12, 8))
        self.bar_frame = bar
        FlatButton(bar, base=ACCENT, hover=ACCENT_HI, text="Начать расчёт",
                   command=self.start_calc, fg="#ffffff", font=(FONT, 11, "bold"),
                   padx=24, pady=10).pack(side="left")
        FlatButton(bar, base=CARD_HI, hover="#2c2c46", text="Пересчитать",
                   command=self.recalc, fg=TEXT, font=(FONT, 10),
                   padx=18, pady=10).pack(side="left", padx=(10, 0))
        FlatButton(bar, base=CARD_HI, hover="#2c2c46", text="Новый расчёт",
                   command=self.new_calc, fg=TEXT, font=(FONT, 10),
                   padx=18, pady=10).pack(side="left", padx=(10, 0))
        tk.Label(bar, text="Все данные обрабатываются локально", bg=BG, fg=FAINT,
                 font=(FONT, 8)).pack(side="right")

        main = tk.Frame(self, bg=BG)
        main.pack(fill="both", expand=True, padx=16)
        self.sf = ScrollFrame(main)
        self.sf.pack(side="left", fill="both", expand=True)

        right = tk.Frame(main, bg=BG, width=330)
        right.pack(side="right", fill="y", padx=(14, 0))
        right.pack_propagate(False)
        self.right_panel = right
        self.btn_export = FlatButton(right, base=GREEN_DK, hover=GREEN_HI, text="Экспорт",
                                     command=self.export_dialog, fg="#ffffff", font=(FONT, 10, "bold"),
                                     padx=14, pady=10)
        self.btn_export.pack(side="bottom", fill="x")
        self.btn_show_total = FlatButton(right, base=ACCENT, hover=ACCENT_HI, text="Показать итог",
                                         command=self.show_breakdown, fg="#ffffff", font=(FONT, 10, "bold"),
                                         padx=14, pady=10)
        self.btn_show_total.pack(side="bottom", fill="x", pady=(14, 0))
        res = tk.Frame(right, bg=CARD, highlightbackground=LINE, highlightthickness=1, padx=18, pady=16)
        res.pack(side="top", fill="x")
        tk.Label(res, text="ИТОГО ПО ОБЪЕКТУ", bg=CARD, fg=MUTED, font=(FONT, 9, "bold")).pack(anchor="w")
        self.total_lbl = tk.Label(res, text="—", bg=CARD, fg=TEXT, font=(FONT, 24, "bold"))
        self.total_lbl.pack(anchor="w", pady=(2, 0))
        self.per_m2_lbl = tk.Label(res, text="Цена за м²: —", bg=CARD, fg=MUTED, font=(FONT, 10))
        self.per_m2_lbl.pack(anchor="w", pady=(2, 6))
        tk.Frame(res, bg=LINE, height=1).pack(fill="x", pady=6)
        self.res_hint = tk.Label(res, text="Загрузите прайс", bg=CARD, fg=AMBER, font=(FONT, 9), anchor="w")
        self.res_hint.pack(fill="x")
        tk.Frame(right, bg=BG).pack(side="bottom", fill="both", expand=True)

    def _on_head_resize(self, event):
        if event.widget is self.head and event.width != self._head_last_w:
            self._head_last_w = event.width
            try:
                self.after_idle(self._check_head_layout)
            except tk.TclError:
                pass

    def _check_head_layout(self):
        try:
            w = self.head.winfo_width()
        except tk.TclError:
            return
        if w <= 1:
            return
        tw = self.head_title_lbl.winfo_reqwidth()
        sw = self.head_sub_lbl.winfo_reqwidth()
        bw = self.head_btns.winfo_reqwidth()
        if min(tw, sw, bw) <= 1:
            self.after(60, self._check_head_layout)
            return
        need = max(tw, sw) + bw + 52
        if not self._head_narrow and w < need:
            self._head_narrow = True
            self.head_btns.grid_forget()
            self.head_btns.grid(row=2, column=0, columnspan=2, sticky="e", pady=(6, 0))
        elif self._head_narrow and w >= need:
            self._head_narrow = False
            self.head_btns.grid_forget()
            self.head_btns.grid(row=0, column=1, rowspan=2, sticky="e")

    def _on_status_resize(self, event):
        if event.widget is self.status and event.width != self._status_last_w:
            self._status_last_w = event.width
            try:
                self.after_idle(self._check_status_layout)
            except tk.TclError:
                pass

    def _check_status_layout(self):
        try:
            w = self.status.winfo_width()
        except tk.TclError:
            return
        if w <= 1:
            return
        avail = max(w - 36, 100)
        fw = tkfont.Font(font=self.file_lbl.cget("font")).measure(self.file_lbl.cget("text")) + 6
        nw = tkfont.Font(font=self.note_lbl.cget("font")).measure(self.note_lbl.cget("text")) + 6
        stacked = fw + nw + 24 > avail
        if stacked:
            self.file_lbl.grid(row=0, column=0, columnspan=2, sticky="w")
            self.note_lbl.grid(row=1, column=0, columnspan=2, sticky="w")
            self.file_lbl.config(wraplength=avail)
            self.note_lbl.config(wraplength=avail)
        else:
            self.file_lbl.grid(row=0, column=0, sticky="w")
            self.note_lbl.grid(row=0, column=1, sticky="e")
            self.file_lbl.config(wraplength=0)
            self.note_lbl.config(wraplength=0)
        self._status_narrow = stacked

    def set_note(self, text, color):
        self.note_lbl.config(text=text, fg=color)
        try:
            self.after_idle(self._check_status_layout)
        except tk.TclError:
            pass

    def _start_dir(self):
        if self._local_dir:
            return self._local_dir
        return os.path.dirname(self.file_path) if self.file_path else os.getcwd()

    def load_dialog(self):
        if self._syncing:
            self.set_note("Идёт обновление с Яндекс.Диска — подождите несколько секунд", MUTED)
            return
        if openpyxl is None:
            messagebox.showerror(
                "Нет библиотеки",
                "Не найдена библиотека openpyxl для чтения .xlsx.\n"
                "Запустите «запуска.bat» или выполните в командной строке:\n\n"
                "python -m pip install openpyxl")
            return
        p = filedialog.askopenfilename(
            initialdir=self._start_dir(),
            filetypes=[("Excel файлы", "*.xlsx"), ("Все файлы", "*.*")],
            title="Выберите файл прайса")
        if p:
            self.load_path(p)

    def load_path(self, path):
        try:
            items, skipped, sheet = parse_price(path)
        except Exception as e:
            messagebox.showerror("Ошибка загрузки", f"Не удалось прочитать файл:\n{e}")
            return
        if not items:
            self.file_path = None
            self.items = []
            self.calculated = False
            self.dirty = False
            self.rebuild_categories()
            self.set_note("В файле нет позиций с ценами — проверьте прайс", AMBER)
            return
        self.file_path = path
        self.items = items
        self.calculated = False
        self.dirty = False
        self.rebuild_categories()
        base = os.path.basename(path)
        note = f"Прайс загружен: {base} · лист «{sheet}» · позиций: {len(items)}"
        if skipped:
            note += f" · пропущено без цены: {skipped}"
        self.set_note(note, GREEN)
        self.file_lbl.config(text=f"Файл: {base}")
        if os.path.abspath(path) != os.path.abspath(YA_TMP_PATH):
            self._local_dir = os.path.dirname(path)
        self.sync_card.pack_forget()
        self.after_idle(self._check_status_layout)

    def sync_yandex(self):
        if openpyxl is None:
            messagebox.showerror(
                "Нет библиотеки",
                "Не найдена библиотека openpyxl для чтения .xlsx.\n"
                "Запустите «запуска.bat» или выполните в командной строке:\n\n"
                "python -m pip install openpyxl")
            return
        if self._syncing:
            self.set_note("Обновление с Яндекс.Диска уже выполняется — подождите", MUTED)
            return
        self._syncing = True
        self._sync_result = None
        self.set_note("Обновление с Яндекс.Диска: получаем актуальный прайс…", MUTED)
        threading.Thread(target=self._sync_worker, daemon=True).start()
        self.after(120, self._sync_poll)

    def _sync_worker(self):
        try:
            path, fname, fdate = fetch_yandex_price(self)
            items, skipped, sheet = parse_price(path)
            if not items:
                raise ValueError("В скачанном с Яндекс.Диска файле нет позиций с ценами")
            result = (None, path, fname, fdate, items, skipped, sheet)
        except Exception as e:
            result = (friendly_net_error(e), None, None, None, None, 0, None)
        self._sync_result = result

    def _sync_poll(self):
        if self._sync_result is None:
            self.after(120, self._sync_poll)
            return
        res = self._sync_result
        self._sync_result = None
        self._syncing = False
        self._sync_apply(*res)

    def _sync_apply(self, err, path, fname, fdate, items, skipped, sheet):
        if err is not None:
            self.set_note("Обновление не выполнено · работает последний загруженный прайс", AMBER)
            messagebox.showerror(
                "Обновление с Яндекс.Диска",
                f"{err}\n\nПриложение продолжает работать с последним загруженным прайсом.")
            return
        old_items = list(self.items)
        added, removed, changed = diff_items(old_items, items)
        self.file_path = path
        self.items = items
        self.calculated = False
        self.dirty = False
        self.rebuild_categories()
        note = f"Прайс обновлён с Яндекс.Диска: {fname} · позиций: {len(items)}"
        if skipped:
            note += f" · пропущено без цены: {skipped}"
        self.set_note(note, GREEN)
        self.file_lbl.config(text=f"Файл: {fname} · Яндекс.Диск · от {_fmt_ya_date(fdate)}")
        self._show_sync_info(old_items, fname, fdate, items, added, removed, changed, sheet)
        self.after_idle(self._check_status_layout)

    def _on_sync_canvas_conf(self, event):
        try:
            self.sync_info_lbl.config(wraplength=max(event.width - 6, 100))
            self.after_idle(self._sync_fit)
        except tk.TclError:
            pass

    def _sync_fit(self):
        try:
            text = self.sync_info_lbl.cget("text")
            if not text:
                return
            w = self.sync_sf.canvas.winfo_width()
        except tk.TclError:
            return
        if w <= 1:
            self.after(60, self._sync_fit)
            return
        f = tkfont.Font(font=self.sync_info_lbl.cget("font"))
        wrap = max(w - 6, 100)
        lines = 0
        for para in text.split("\n"):
            lines += max(1, int(-(-f.measure(para) // wrap)))
        need = int(lines * f.metrics("linespace")) + 6
        cap = 72
        try:
            self.sync_body.config(height=min(need, cap))
            if need > cap:
                if not self.sync_sf.vsb.winfo_manager():
                    self.sync_sf.vsb.pack(side="right", fill="y")
            else:
                if self.sync_sf.vsb.winfo_manager():
                    self.sync_sf.vsb.pack_forget()
        except tk.TclError:
            pass

    def _show_sync_info(self, old_items, fname, fdate, items, added, removed, changed, sheet):
        first = f"Файл: {fname} · на Диске: {_fmt_ya_date(fdate)} · позиций: {len(items)}"
        if sheet:
            first += f" · лист «{sheet}»"
        lines = [first]
        if not old_items:
            lines.append("Изменения: первый импорт с Яндекс.Диска")
        else:
            parts = []
            if added:
                parts.append(f"новых позиций: {len(added)}")
            if removed:
                parts.append(f"удалено позиций: {len(removed)}")
            if changed:
                parts.append(f"изменилось цен: {len(changed)}")
            lines.append("Изменения: " + ("; ".join(parts) if parts else "нет — состав и цены совпадают"))
            exs = []
            for o, n in changed[:4]:
                exs.append(f"«{o.name}»: {money(o.price)} → {money(n.price)}")
            for it in removed[:3]:
                exs.append(f"удалена «{it.name}»")
            for it in added[:3]:
                exs.append(f"добавлена «{it.name}»")
            exs = exs[:6]
            if exs:
                lines.append("; ".join(exs))
        self.sync_info_lbl.config(text="\n".join(lines))
        self.sync_card.pack(fill="x", padx=16, before=self.area_card)
        try:
            self.after_idle(self._sync_fit)
        except tk.TclError:
            pass

    def rebuild_categories(self):
        for w in self.sf.inner.winfo_children():
            w.destroy()
        self.blocks = {}
        if not self.items:
            ph = tk.Frame(self.sf.inner, bg=BG)
            tk.Label(ph, text="Загрузите прайс", bg=BG, fg=MUTED,
                     font=(FONT, 16, "bold")).pack(pady=(90, 6))
            tk.Label(ph, text="Нажмите «Обновить данные» (актуальный прайс с Яндекс.Диска) или «Загрузить из файла»\n"
                              "Нужен файл .xlsx с листом «Прайс» (колонки: Категория, Позиция, Ед. изм., Цена, Комментарий)",
                     bg=BG, fg=FAINT, font=(FONT, 9), justify="center").pack()
            ph.pack(fill="x")
            self.total_lbl.config(text="—")
            self.per_m2_lbl.config(text="Цена за м²: —")
            self.res_hint.config(text="Загрузите прайс", fg=AMBER)
            return
        order = []
        seen = set()
        for it in self.items:
            if it.cat not in seen:
                seen.add(it.cat)
                order.append(it.cat)
        for cat in order:
            cat_items = [it for it in self.items if it.cat == cat]
            kind = "single" if cat.strip().lower() in SINGLE_CATEGORIES else "multi"
            self._category_card(cat, kind, cat_items)
        self._shown = 0
        self.total_lbl.config(text="—")
        self.per_m2_lbl.config(text="Цена за м²: —")
        self.res_hint.config(text="Нажмите «Начать расчёт»", fg=ACCENT_HI)

    def _category_card(self, cat, kind, cat_items):
        card = tk.Frame(self.sf.inner, bg=CARD, highlightbackground=LINE,
                        highlightthickness=1, padx=14, pady=12)
        card.pack(fill="x", pady=(0, 12))
        head = tk.Frame(card, bg=CARD)
        head.pack(fill="x", pady=(0, 8))
        tk.Label(head, text=cat, bg=CARD, fg=TEXT, font=(FONT, 12, "bold")).pack(side="left")
        tk.Label(head, text="выберите один вариант" if kind == "single" else "включите нужные позиции",
                 bg=CARD, fg=MUTED, font=(FONT, 8)).pack(side="right")
        rows = []
        for it in cat_items:
            row = OptionRow(card, kind, it, self._on_option)
            row.pack(fill="x", pady=1)
            rows.append(row)
        if kind == "single":
            rows[0].set_selected(True, animate=False)
        self.blocks[cat] = {"kind": kind, "rows": rows}

    def _on_option(self, row):
        if row.mode == "single":
            blk = self.blocks.get(row.item.cat)
            if blk:
                for r in blk["rows"]:
                    r.set_selected(r is row)
        self._mark_dirty()

    def _mark_dirty(self):
        if self.calculated and not self.dirty:
            self.dirty = True
            self.set_note("Есть изменения — нажмите «Пересчитать»", AMBER)

    def _on_area(self, value):
        try:
            self.area_val.config(text=f"{int(float(value))} м²")
        except tk.TclError:
            return
        self._mark_dirty()

    def compute(self):
        area = int(self.area_var.get())
        rows = []
        for cat, blk in self.blocks.items():
            for r in blk["rows"]:
                if r.selected:
                    it = r.item
                    cost = it.price * area if it.area_based else it.price
                    rows.append((cat, it, cost))
        return sum(c for _, _, c in rows), rows, area

    def start_calc(self):
        if not self.items:
            self.set_note("Загрузите прайс — «Обновить данные» или «Загрузить из файла»", AMBER)
            return
        self._do_calc()

    def recalc(self):
        if not self.items:
            self.set_note("Загрузите прайс — «Обновить данные» или «Загрузить из файла»", AMBER)
            return
        self._do_calc()

    def _do_calc(self):
        total, rows, area = self.compute()
        self.last_total = total
        self.last_rows = rows
        self.last_area = area
        self.calculated = True
        self.dirty = False
        self.animate_total(total)
        self.per_m2_lbl.config(text=f"Цена за м²: {money(total / area)}")
        self.res_hint.config(text=f"Площадь: {area} м² · позиций в расчёте: {len(rows)}", fg=GREEN)
        self.set_note(f"Расчёт выполнен · площадь {area} м² · сумма {money(total)}", GREEN)

    def new_calc(self):
        if not self.items:
            self.set_note("Загрузите прайс — «Обновить данные» или «Загрузить из файла»", AMBER)
            return
        self.calculated = False
        self.dirty = False
        self.area_var.set(100)
        for cat, blk in self.blocks.items():
            first = True
            for r in blk["rows"]:
                r.set_selected(first and blk["kind"] == "single", animate=False)
                first = False
        self._shown = 0
        self.total_lbl.config(text="—")
        self.per_m2_lbl.config(text="Цена за м²: —")
        self.res_hint.config(text="Нажмите «Начать расчёт»", fg=ACCENT_HI)
        self.set_note("Новый расчёт: настройки сброшены", MUTED)

    def animate_total(self, target):
        if self._total_job:
            try:
                self.after_cancel(self._total_job)
            except tk.TclError:
                return
            self._total_job = None
        start = self._shown
        steps = 16
        state = [0]

        def tick():
            try:
                state[0] += 1
                t = state[0] / steps
                v = start + (target - start) * t
                self.total_lbl.config(text=money(v))
                if state[0] < steps:
                    self._total_job = self.after(16, tick)
                else:
                    self._shown = target
                    self.total_lbl.config(text=money(target))
            except tk.TclError:
                return

        tick()

    def show_breakdown(self):
        if not self.calculated:
            messagebox.showinfo("Итог", "Сначала нажмите «Начать расчёт».")
            return
        win = tk.Toplevel(self)
        win.title("Итог расчёта")
        win.configure(bg=BG)
        win.geometry("640x700")
        win.transient(self)
        head = tk.Frame(win, bg=CARD, highlightbackground=LINE, highlightthickness=1, padx=18, pady=14)
        head.pack(fill="x", padx=14, pady=(14, 0))
        tk.Label(head, text="ИТОГО ПО ОБЪЕКТУ", bg=CARD, fg=MUTED, font=(FONT, 9, "bold")).pack(anchor="w")
        tk.Label(head, text=money(self.last_total), bg=CARD, fg=TEXT, font=(FONT, 20, "bold")).pack(anchor="w")
        tk.Label(head, text=f"Цена за м²: {money(self.last_total / self.last_area)} · площадь: {self.last_area} м²",
                 bg=CARD, fg=MUTED, font=(FONT, 9)).pack(anchor="w")
        sf = ScrollFrame(win)
        sf.pack(fill="both", expand=True, padx=14, pady=10)
        for cat, grp in groupby(self.last_rows, key=lambda r: r[0]):
            card = tk.Frame(sf.inner, bg=CARD, highlightbackground=LINE, highlightthickness=1, padx=12, pady=8)
            card.pack(fill="x", pady=(0, 8))
            tk.Label(card, text=cat, bg=CARD, fg=ACCENT_HI, font=(FONT, 10, "bold")).pack(anchor="w", pady=(0, 4))
            for _, it, cost in grp:
                line = tk.Frame(card, bg=CARD)
                line.pack(fill="x")
                tk.Label(line, text=it.name, bg=CARD, fg=TEXT, font=(FONT, 9), anchor="w",
                         wraplength=330, justify="left").pack(side="left", fill="x", expand=True)
                tk.Label(line, text=money(cost), bg=CARD, fg=TEXT, font=(FONT, 9, "bold")).pack(side="right")
                unit_disp = it.unit or ("м²" if it.area_based else "шт")
                tk.Label(card, text=f"{unit_disp} · {money(it.price)}" + ("/м²" if it.area_based else ""),
                         bg=CARD, fg=FAINT, font=(FONT, 7)).pack(anchor="w")
        btn = tk.Frame(win, bg=BG)
        btn.pack(fill="x", padx=14, pady=(0, 14))
        FlatButton(btn, base=CARD_HI, hover="#2c2c46", text="Закрыть",
                   command=win.destroy, fg=TEXT, font=(FONT, 10), padx=18, pady=8).pack(side="right")

    def export_dialog(self):
        if openpyxl is None:
            messagebox.showerror("Нет библиотеки", "Для экспорта нужна библиотека openpyxl.\nЗапустите «запуска.bat».")
            return
        if not self.calculated:
            messagebox.showinfo("Экспорт", "Сначала выполните расчёт.")
            return
        p = filedialog.asksaveasfilename(
            initialdir=self._start_dir(),
            initialfile="Расчёт_каркасного_дома.xlsx",
            defaultextension=".xlsx",
            filetypes=[("Excel файлы", "*.xlsx")],
            title="Сохранить расчёт")
        if not p:
            return
        try:
            self.export_to(p)
        except Exception as e:
            messagebox.showerror("Ошибка экспорта", f"Не удалось сохранить файл:\n{e}")
            return
        messagebox.showinfo("Экспорт", f"Файл сохранён:\n{p}")

    def export_to(self, path):
        wb = Workbook()
        ws = wb.active
        ws.title = "Расчёт"
        bold = XlFont(bold=True)
        ws.append(["Расчёт стоимости каркасного дома"])
        ws.append(["Дата:", datetime.date.today().strftime("%d.%m.%Y")])
        ws.append(["Площадь дома:", self.last_area, "м²"])
        ws.append([])
        ws.append(["Категория", "Позиция", "Ед. изм.", "Цена", "Стоимость"])
        for cell in ws[5]:
            cell.font = bold
        for cat, it, cost in self.last_rows:
            unit_disp = it.unit or ("м²" if it.area_based else "шт")
            ws.append([cat, it.name, unit_disp, it.price, round(cost, 2)])
        ws.append([])
        ws.append(["ИТОГО", "", "", "", round(self.last_total, 2)])
        ws.append(["Цена за м²", "", "", "", round(self.last_total / self.last_area, 2)])
        for r in (ws.max_row - 1, ws.max_row):
            for cell in ws[r]:
                cell.font = bold
        for i, w in enumerate([24, 44, 10, 14, 16], 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        wb.save(path)


def main():
    if sys.platform == "win32":
        try:
            ctypes.windll.shcore.SetProcessDpiAwareness(1)
        except Exception:
            pass
    App().mainloop()


if __name__ == "__main__":
    main()
